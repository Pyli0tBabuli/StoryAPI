"""Build the script reference document out of the dump the game wrote.

The dump itself is a plain text file with everything the game's script contains:
every line of dialogue with the number that identifies it, and every step of every
block. That is exactly what a mod author needs in order to hook a scene, but as a
text file it is awkward to search and impossible to print.

Run:  python build_script_reference.py <data folder> <output folder>
"""

import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def parse_reference(path):
    """Split the dump into the dialogue part and the block part."""
    lines = []
    with open(path, encoding="utf-8", errors="replace") as handle:
        lines = handle.read().split("\n")

    sections = []          # (kind, title, [lines])
    current = None
    in_blocks = False

    for line in lines:
        if line.startswith("# Blocks"):
            in_blocks = True
            if current:
                current = None
            continue

        if line.startswith("## "):
            title = line[3:].rstrip()
            match = re.match(r"^(.*?)\s+\((\d+)\s+(line|step)\(s\)\)$", title)
            if match:
                name, count, unit = match.group(1), int(match.group(2)), match.group(3)
                current = (unit, name, [])
                sections.append(current)
            else:
                current = None
            continue

        if current is not None and line.strip():
            current[2].append(line.rstrip())

    return sections


def mono(document, text, size=7, bold=False, color=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return paragraph


def build(data_dir, out_dir):
    source = os.path.join(data_dir, "script-reference.txt")
    sections = parse_reference(source)

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)

    heading = document.add_heading("DDLC Plus - Script reference", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    dialogue = [s for s in sections if s[0] == "line"]
    blocks = [s for s in sections if s[0] == "step"]

    total_lines = sum(len(s[2]) for s in dialogue)
    total_steps = sum(len(s[2]) for s in blocks)

    document.add_paragraph(
        "Everything the game's script contains, taken from the running game itself.\n"
        "\n"
        "Part one lists every line of dialogue. A line is identified by the pair of a\n"
        "label and a number, and that pair is what a mod needs in order to replace a\n"
        "line or to continue a scene after it:\n"
        "\n"
        "    override ch1_m_start = 1271520166 \"My own line.\"\n"
        "    hook ch1_m_start after 1271520166:\n"
        "\n"
        "Part two lists every block of the script and the steps it is made of. A step\n"
        "is one instruction: a scene, a show, a line of dialogue, a jump. The text in\n"
        "quotes is the command as the game wrote it, so it names the background, the\n"
        "image and the transition the game itself uses."
    )

    document.add_paragraph(
        f"{len(dialogue)} label(s), {total_lines} line(s), "
        f"{len(blocks)} block(s), {total_steps} step(s)."
    )

    document.add_page_break()

    document.add_heading("Part 1 - Dialogue", level=1)
    mono(document, "label = line number  text", 8, bold=True,
         color=RGBColor(0x66, 0x66, 0x66))

    for _, name, rows in dialogue:
        document.add_heading(name, level=2)
        for row in rows:
            mono(document, row)

    document.add_page_break()

    document.add_heading("Part 2 - Blocks and steps", level=1)
    mono(document, "step  type  command", 8, bold=True,
         color=RGBColor(0x66, 0x66, 0x66))

    for _, name, rows in blocks:
        document.add_heading(name, level=2)
        for row in rows:
            mono(document, row)

    out = os.path.join(out_dir, "07-Script-Reference.docx")
    document.save(out)
    return out, len(dialogue), total_lines, len(blocks), total_steps


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1

    data, out_dir = sys.argv[1], sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)

    out, labels, lines, blocks, steps = build(data, out_dir)
    print(f"  {os.path.basename(out):<34} {labels} label(s), {lines} line(s), "
          f"{blocks} block(s), {steps} step(s)  "
          f"{os.path.getsize(out) / 1048576:.1f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
